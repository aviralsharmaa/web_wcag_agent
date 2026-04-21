"""Convert LICMF Investor App WCAG report to ICICI-format XLSX with OneDrive screenshots."""
import os
import sys
import time
import json
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Configuration ──
BASE = "reports/com-licmf-investor__v1-0-5__20260326_174909/checklist_reports"
SRC_FILE = "LICMF_Investor_App_WCAG_Audit_Report (2).xlsx"
OUT_FILE = "LICMF_Investor_App_WCAG_Report.xlsx"

CHECKLIST_TO_FOLDER = {
    'Labels or Instructions': '01_labels_or_instructions',
    'Identify Input Purpose': '02_identify_input_purpose',
    'Name, Role, Value': '03_name_role_value',
    'Status Messages': '04_status_messages',
    'Keyboard': '05_keyboard',
    'Focus Order': '06_focus_order',
    'Meaningful Sequence': '07_meaningful_sequence',
    'Error Identification': '08_error_identification',
    'Error Prevention (Legal, Financial, Data)': '09_error_prevention_legal_financial_data',
    'Error Suggestion': '10_error_suggestion',
    'Label in Name': '11_label_in_name',
}

ISSUE_TYPE_MAP = {
    'missing_accessible_name': 'Screen Reader',
    'input_missing_label': 'Screen Reader',
    'naf_element': 'Screen Reader',
    'not_focusable': 'Keyboard Navigation',
    'small_touch_target': 'Other A11y',
    'focus_order_issue': 'Keyboard Navigation',
    'missing_content_description': 'Screen Reader',
    'error_not_identified': 'Screen Reader',
    'status_not_announced': 'Screen Reader',
    'input_purpose_missing': 'Screen Reader',
    'label_mismatch': 'Screen Reader',
}

SEVERITY_DISPLAY = {
    'Critical': 'Blocker',
    'Major': 'High',
    'Moderate': 'Medium',
    'Minor': 'Low',
}

from src.accessibility_scanner.xlsx_report import (
    WHAT_TO_FIX, HOW_TO_FIX, EXPECTED_RESULT, CODE_FIX,
)

# ── Styling ──
_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_HEADER_FILL = PatternFill("solid", fgColor="4A3728")
_SECTION_FILL = PatternFill("solid", fgColor="B87333")
_SALMON_FILL = PatternFill("solid", fgColor="E6B8AF")
_GREEN_FILL = PatternFill("solid", fgColor="00FF00")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DARK_FONT = Font(name="Calibri", size=11, bold=True, color="2E3436")
_CODE_FIX_FILL = PatternFill("solid", fgColor="F2F2F2")
_CODE_FIX_FONT = Font(name="Consolas", size=9, color="333333")
_SEV_FILLS = {
    "Blocker": (PatternFill("solid", fgColor="FF0000"), Font(name="Calibri", size=10, bold=True, color="FFFFFF")),
    "High":    (PatternFill("solid", fgColor="FF6600"), Font(name="Calibri", size=10, bold=True, color="FFFFFF")),
    "Medium":  (PatternFill("solid", fgColor="FFCC00"), Font(name="Calibri", size=10, bold=True, color="000000")),
    "Low":     (PatternFill("solid", fgColor="99CCFF"), Font(name="Calibri", size=10, bold=True, color="000000")),
}


def find_screenshot(checklist_name, issue_id, ann_filename):
    """Find the full-page annotated screenshot for an issue."""
    folder = CHECKLIST_TO_FOLDER.get(checklist_name)
    if not folder:
        return None
    # Prefer full-page annotated screenshot (shows entire screen with issue highlighted)
    if ann_filename:
        ann_path = os.path.join(BASE, folder, 'states', 'annotated', ann_filename)
        if os.path.isfile(ann_path):
            return ann_path
    # Fallback to crop only if full-page not available
    crop_path = os.path.join(BASE, folder, 'issues', f'{issue_id}_crop.png')
    if os.path.isfile(crop_path):
        return crop_path
    return None


def label_value(ws, row, label, value):
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = Font(name="Calibri", size=11)
    lc.border = _THIN
    vc = ws.cell(row=row, column=3, value=value)
    vc.font = Font(name="Calibri", size=11)
    vc.border = _THIN
    vc.alignment = Alignment(wrap_text=True)


def main():
    sys.stdout.reconfigure(line_buffering=True)

    # ── Step 1: Read source (only needed columns, skip XML dumps) ──
    print("Reading source report (skipping XML columns)...")
    src_wb = openpyxl.load_workbook(SRC_FILE, data_only=True, read_only=True)

    sd_ws = src_wb['Screen Details']
    screen_info = {}
    for row_idx, row in enumerate(sd_ws.iter_rows(min_row=4, values_only=True)):
        if not row[0]:  # column A = #
            continue
        screen_info[row[1]] = {
            'annotated_screenshot': row[2],
            'activity': row[3],
            'top_violations': row[11] if len(row) > 11 else '',
        }
    print(f"Screens: {len(screen_info)}")

    fd_ws = src_wb['Failure Details']
    issues = []
    for row_idx, row in enumerate(fd_ws.iter_rows(min_row=3, values_only=True)):
        if not row[0]:  # column A = #
            continue
        # Only read columns A-L (0-11), skip M-O (XML dumps)
        issues.append({
            'num': row[0],
            'checklist': row[1] or '',
            'issue_id': row[2] or '',
            'screen': row[3] or '',
            'issue_type': row[4] or '',
            'wcag_nums': row[5] or '',
            'checkpoint_name': row[6] or '',
            'level': row[7] or 'A',
            'rationale': row[8] or '',
            'severity': row[9] or 'Moderate',
            'fix': row[10] or '',
            'annotated_screenshot': row[11] or '',
        })
    src_wb.close()
    print(f"Total issues: {len(issues)}")

    # ── Step 2: Map screenshots ──
    print("Mapping screenshots to local files...")
    for issue in issues:
        issue['local_screenshot'] = find_screenshot(
            issue['checklist'], issue['issue_id'], issue['annotated_screenshot']
        )
    with_ss = sum(1 for i in issues if i['local_screenshot'])
    print(f"Issues with screenshots: {with_ss}/{len(issues)}")

    # ── Step 3: Upload to OneDrive ──
    print("Uploading screenshots to OneDrive...")
    from src.accessibility_scanner.image_uploader import OneDriveUploader
    uploader = OneDriveUploader()
    upload_cache = {}
    upload_count = 0
    upload_errors = 0

    unique_paths = sorted(set(i['local_screenshot'] for i in issues if i['local_screenshot']))
    total = len(unique_paths)
    print(f"Unique screenshots to upload: {total}")

    for path in unique_paths:
        for attempt in range(3):
            try:
                url = uploader.upload(path)
                upload_cache[path] = url
                upload_count += 1
                if upload_count % 25 == 0:
                    print(f"  Progress: {upload_count}/{total} ({100*upload_count//total}%)")
                time.sleep(0.15)
                break
            except Exception as e:
                if attempt == 2:
                    print(f"  FAILED: {os.path.basename(path)} - {e}")
                    upload_cache[path] = ""
                    upload_errors += 1
                else:
                    time.sleep(2 ** attempt)

    print(f"Upload complete: {upload_count} succeeded, {upload_errors} failed")

    # Map URLs back
    for issue in issues:
        path = issue.get('local_screenshot')
        issue['screenshot_url'] = upload_cache.get(path, '') if path else ''

    # ── Step 4: Generate report ──
    print("Generating ICICI-format report...")
    wb = Workbook()

    # === COVER SHEET ===
    ws = wb.active
    ws.title = "Cover"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 25

    ws.merge_cells("A1:D3")
    c = ws.cell(row=1, column=1, value="Accessibility Audit Report")
    c.font = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
    c.fill = _HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(1, 4):
        ws.row_dimensions[r].height = 28
        for col in range(1, 5):
            ws.cell(row=r, column=col).fill = _HEADER_FILL

    ws.merge_cells("A5:D5")
    c5 = ws.cell(row=5, column=1, value="ASSET DETAILS")
    c5.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c5.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=5, column=col).fill = _SECTION_FILL

    label_value(ws, 6, "Asset / Product Name", "LICMF Investor App")
    label_value(ws, 7, "App Version", "v1.0.5 (Build 10006)")
    label_value(ws, 8, "Standard", "WCAG 2.1 A and AA")
    label_value(ws, 9, "Scan Mode", "Android Mobile App")
    label_value(ws, 10, "Screens Analyzed", str(len(screen_info)))

    ws.merge_cells("A12:D12")
    c12 = ws.cell(row=12, column=1, value="SUBMISSION DETAILS")
    c12.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c12.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=12, column=col).fill = _SECTION_FILL
    label_value(ws, 13, "Date of Submission", "14 Apr 2026")

    ws.merge_cells("A15:D15")
    c15 = ws.cell(row=15, column=1, value="ABOUT THIS AUDIT")
    c15.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c15.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=15, column=col).fill = _SECTION_FILL

    ws.merge_cells("A16:D16")
    c16 = ws.cell(row=16, column=1, value=(
        "This report presents findings from an accessibility audit conducted in accordance with the Web Content "
        "Accessibility Guidelines (WCAG) 2.1. Issues are categorised by severity (Blocker / High / Medium / Low) "
        "and mapped to the relevant WCAG success criteria. Conformance levels A and AA are referenced throughout."
    ))
    c16.font = Font(name="Calibri", size=10)
    c16.alignment = Alignment(wrap_text=True, vertical="top")
    c16.border = _THIN
    ws.row_dimensions[16].height = 72

    ws.merge_cells("A18:D18")
    c18 = ws.cell(row=18, column=1, value="AUDITOR DETAILS")
    c18.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c18.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=18, column=col).fill = _SECTION_FILL
    label_value(ws, 19, "Auditor", "Automated WCAG Scanner")
    label_value(ws, 20, "Organisation", "OnFinance AI")

    ws.merge_cells("A22:D22")
    c22 = ws.cell(row=22, column=1, value="CONFIDENTIAL \u2014 For client use only")
    c22.font = Font(name="Calibri", size=11, italic=True)
    c22.alignment = Alignment(horizontal="center", vertical="center")
    c22.border = _THIN

    # === STATUS SHEET ===
    ws_s = wb.create_sheet("Status")
    for col_letter, w in zip("ABCDEFGH", [30, 45, 35, 18, 18, 18, 18, 18]):
        ws_s.column_dimensions[col_letter].width = w

    for col_idx, hdr in enumerate([
        "Screen Label", "Activity", "Notes", "Labels/Instructions",
        "Keyboard", "Focus Order", "Name, Role, Value", "Touch Targets",
    ], start=1):
        c = ws_s.cell(row=1, column=col_idx, value=hdr)
        c.fill = _SALMON_FILL
        c.font = _DARK_FONT
        c.alignment = _CENTER
        c.border = _THIN
    ws_s.freeze_panes = "A2"

    screen_checklists = {}
    screen_issue_types = {}
    for issue in issues:
        scr = issue['screen']
        screen_checklists.setdefault(scr, set()).add(issue['checklist'])
        screen_issue_types.setdefault(scr, set()).add(issue['issue_type'])

    for screen_label in sorted(screen_info.keys()):
        info = screen_info[screen_label]
        row_idx = ws_s.max_row + 1
        failed = screen_checklists.get(screen_label, set())

        for col_idx, val in enumerate([screen_label, info.get('activity', ''), ""], start=1):
            c = ws_s.cell(row=row_idx, column=col_idx, value=val)
            c.border = _THIN
            c.alignment = _WRAP_TOP

        for cl_name, col in [('Labels or Instructions', 4), ('Keyboard', 5),
                             ('Focus Order', 6), ('Name, Role, Value', 7)]:
            c = ws_s.cell(row=row_idx, column=col)
            if cl_name in failed:
                c.value = "Issues Found"
                c.fill = PatternFill("solid", fgColor="FFA500")
            else:
                c.value = "Done"
                c.fill = _GREEN_FILL
            c.border = _THIN
            c.alignment = _CENTER

        has_touch = 'small_touch_target' in screen_issue_types.get(screen_label, set())
        c8 = ws_s.cell(row=row_idx, column=8)
        c8.value = "Issues Found" if has_touch else "Done"
        c8.fill = PatternFill("solid", fgColor="FFA500") if has_touch else _GREEN_FILL
        c8.border = _THIN
        c8.alignment = _CENTER

    # === ISSUES SHEET ===
    ws_i = wb.create_sheet("Issues")
    col_widths = [30, 22, 18, 20, 12, 35, 45, 45, 20, 22, 45, 55, 14, 22, 12, 18]
    for i, w in enumerate(col_widths):
        ws_i.column_dimensions[get_column_letter(i + 1)].width = w

    headers_i = [
        "Issue Title", "Page Title", "Type", "AT/Browser", "Severity",
        "Action Performed", "Actual Result", "Expected Result",
        "Failed WCAG 2.1 checkpoint(s)", "Screencast",
        "Suggested Resolutions", "Suggested Code Fix",
        "Dev Status", "Dev Comments", "QA Status", "QA Comments",
    ]
    for col_idx, hdr in enumerate(headers_i, start=1):
        c = ws_i.cell(row=1, column=col_idx, value=hdr)
        c.fill = _SALMON_FILL
        c.font = _DARK_FONT
        c.alignment = _CENTER
        c.border = _THIN
    ws_i.freeze_panes = "A2"

    action_performed = (
        "1) Launch the LICMF Investor App (v1.0.5).\n"
        "2) Navigate to the target screen.\n"
        "3) Run automated WCAG 2.1 accessibility scan."
    )

    for idx, issue in enumerate(issues):
        it = ISSUE_TYPE_MAP.get(issue['issue_type'], 'Screen Reader')
        sev = SEVERITY_DISPLAY.get(issue['severity'], 'Medium')
        primary_wcag = issue['wcag_nums'].split(',')[0].strip() if issue['wcag_nums'] else ''

        expected = EXPECTED_RESULT.get(primary_wcag,
            f"Content must conform to WCAG 2.1 {primary_wcag} {issue['checkpoint_name']} (Level {issue['level']}).")

        what = WHAT_TO_FIX.get(primary_wcag, issue['fix'][:200] if issue['fix'] else '')
        how = HOW_TO_FIX.get(primary_wcag, '')
        suggested = (what + "\n\n" + "\n".join(how.split("\n")[:2])) if how else (issue['fix'] or what)

        code_fix = CODE_FIX.get(primary_wcag, '')
        title_summary = issue['rationale'][:80] if issue['rationale'] else issue['issue_type']
        issue_title = f"{it}: {issue['screen']}: {title_summary}"
        wcag_str = f"{issue['wcag_nums']} {issue['checkpoint_name']} (Level {issue['level']})"

        row_idx = ws_i.max_row + 1
        ws_i.row_dimensions[row_idx].height = 80

        row_values = [
            issue_title, issue['screen'], it,
            "Android (Automated WCAG Scanner)", sev, action_performed,
            issue['rationale'], expected, wcag_str,
            "", suggested, code_fix, "", "", "", "",
        ]
        for col_idx, val in enumerate(row_values, start=1):
            c = ws_i.cell(row=row_idx, column=col_idx, value=val)
            c.border = _THIN
            c.alignment = _WRAP_TOP

        sev_style = _SEV_FILLS.get(sev)
        if sev_style:
            ws_i.cell(row=row_idx, column=5).fill = sev_style[0]
            ws_i.cell(row=row_idx, column=5).font = sev_style[1]

        sc_cell = ws_i.cell(row=row_idx, column=10)
        if issue.get('screenshot_url'):
            sc_cell.value = "View Screenshot"
            sc_cell.hyperlink = issue['screenshot_url']
            sc_cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            sc_cell.alignment = Alignment(horizontal="center", vertical="center")
        elif issue.get('annotated_screenshot'):
            sc_cell.value = issue['annotated_screenshot']
            sc_cell.alignment = _WRAP_TOP

        if code_fix:
            code_cell = ws_i.cell(row=row_idx, column=12)
            code_cell.font = _CODE_FIX_FONT
            code_cell.fill = _CODE_FIX_FILL
            code_cell.alignment = Alignment(wrap_text=True, vertical="top")

    print(f"Issues written: {len(issues)}")
    wb.save(OUT_FILE)
    print(f"\nReport saved: {OUT_FILE}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
