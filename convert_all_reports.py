"""Convert WCAG audit reports to ICICI-format XLSX with OneDrive screenshots.

Handles two source formats:
  Format A (Kotak Neo, Kinsite, LICMF): Failure Details has Checklist column
  Format B (HSL InvestRight): No Checklist column, fewer Screen Details columns
"""
import os
import sys
import time
import json
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

from src.accessibility_scanner.xlsx_report import (
    WHAT_TO_FIX, HOW_TO_FIX, EXPECTED_RESULT, CODE_FIX,
)

# ── App configurations ──
APPS = [
    {
        "name": "Kotak Neo",
        "version": "v2.2.73",
        "base": "reports/com-kotak-neo__v2-2-73__20260408_130507/checklist_reports",
        "src_file": "Kotak_Neo_WCAG_Audit_Report (2).xlsx",
        "out_file": "Kotak_Neo_WCAG_Report.xlsx",
        "scan_mode": "Android Mobile App",
        "format": "A",  # Has Checklist column in Failure Details
    },
    {
        "name": "HSL InvestRight",
        "version": "v4.7.2 (UAT)",
        "base": "reports/com-hsl-investright-uat__v4-7-2__20260327_174224/checklist_reports",
        "src_file": "com.hsl.investright.uat_WCAG_Audit_Report.xlsx",
        "out_file": "HSL_InvestRight_WCAG_Report.xlsx",
        "scan_mode": "Android Mobile App",
        "format": "B",  # No Checklist column
    },
    {
        "name": "KIE Research (Kinsite)",
        "version": "v1.1.17",
        "base": "reports/com-kotak-kinsite-apps-kinsite-research-kinsite__v1-1-17__20260401_214600/checklist_reports",
        "src_file": "KIE_Research_WCAG_Audit_Report (3).xlsx",
        "out_file": "KIE_Research_WCAG_Report.xlsx",
        "scan_mode": "Android Mobile App",
        "format": "A",
    },
]

ISSUE_TYPE_MAP = {
    'missing_accessible_name': 'Screen Reader',
    'input_missing_label': 'Screen Reader',
    'naf_element': 'Screen Reader',
    'not_focusable': 'Keyboard Navigation',
    'small_touch_target': 'Other A11y',
    'focus_order_issue': 'Keyboard Navigation',
    'missing_content_description': 'Screen Reader',
    'error_not_identified': 'Screen Reader',
    'status_not_announced': 'Screen Reader',
    'input_purpose_missing': 'Screen Reader',
    'label_mismatch': 'Screen Reader',
    'decorative_image_exposed': 'Screen Reader',
    'heading_structure_issue': 'Screen Reader',
    'color_contrast_fail': 'Color Contrast',
    'orientation_locked': 'Other A11y',
    'language_missing': 'Screen Reader',
    'reflow_issue': 'Zoom',
    'text_spacing_issue': 'Zoom',
    'non_text_contrast_fail': 'Color Contrast',
}

SEVERITY_DISPLAY = {
    'Critical': 'Blocker',
    'Major': 'High',
    'Moderate': 'Medium',
    'Minor': 'Low',
}

# ── Styling ──
_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_HEADER_FILL = PatternFill("solid", fgColor="4A3728")
_SECTION_FILL = PatternFill("solid", fgColor="B87333")
_SALMON_FILL = PatternFill("solid", fgColor="E6B8AF")
_GREEN_FILL = PatternFill("solid", fgColor="00FF00")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DARK_FONT = Font(name="Calibri", size=11, bold=True, color="2E3436")
_CODE_FIX_FILL = PatternFill("solid", fgColor="F2F2F2")
_CODE_FIX_FONT = Font(name="Consolas", size=9, color="333333")
_SEV_FILLS = {
    "Blocker": (PatternFill("solid", fgColor="FF0000"), Font(name="Calibri", size=10, bold=True, color="FFFFFF")),
    "High":    (PatternFill("solid", fgColor="FF6600"), Font(name="Calibri", size=10, bold=True, color="FFFFFF")),
    "Medium":  (PatternFill("solid", fgColor="FFCC00"), Font(name="Calibri", size=10, bold=True, color="000000")),
    "Low":     (PatternFill("solid", fgColor="99CCFF"), Font(name="Calibri", size=10, bold=True, color="000000")),
}


def build_screenshot_index(base_path):
    """Build filename→full_path index for all annotated screenshots."""
    index = {}
    base = Path(base_path)
    if not base.exists():
        return index
    for ann_dir in base.glob("*/states/annotated"):
        for png in ann_dir.glob("*_annotated.png"):
            index[png.name] = str(png)
    print(f"  Screenshot index: {len(index)} files")
    return index


def read_source_format_a(src_file):
    """Read Format A: has Checklist column (Kotak Neo, Kinsite, LICMF)."""
    print(f"  Reading Format A source: {src_file}")
    wb = openpyxl.load_workbook(src_file, data_only=True, read_only=True)

    sd_ws = wb['Screen Details']
    screen_info = {}
    for row in sd_ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        screen_info[row[1]] = {
            'annotated_screenshot': row[2] if len(row) > 2 else '',
            'activity': row[3] if len(row) > 3 else '',
            'top_violations': row[11] if len(row) > 11 else '',
        }

    fd_ws = wb['Failure Details']
    issues = []
    for row in fd_ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        issues.append({
            'num': row[0],
            'checklist': row[1] or '',
            'issue_id': row[2] or '',
            'screen': row[3] or '',
            'issue_type': row[4] or '',
            'wcag_nums': row[5] or '',
            'checkpoint_name': row[6] or '',
            'level': row[7] or 'A',
            'rationale': row[8] or '',
            'severity': row[9] or 'Moderate',
            'fix': row[10] or '',
            'annotated_screenshot': row[11] or '',
        })
    wb.close()
    return screen_info, issues


def read_source_format_b(src_file):
    """Read Format B: no Checklist column (HSL InvestRight)."""
    print(f"  Reading Format B source: {src_file}")
    wb = openpyxl.load_workbook(src_file, data_only=True, read_only=True)

    sd_ws = wb['Screen Details']
    screen_info = {}
    for row in sd_ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        screen_info[row[1]] = {
            'annotated_screenshot': row[2] if len(row) > 2 else '',
            'activity': row[3] if len(row) > 3 else '',
            'top_violations': '',
        }

    fd_ws = wb['Failure Details']
    issues = []
    # Format B columns: #, Issue ID, Screen, Issue Type, WCAG #, Checkpoint Name, Level, Failure Rationale, Severity, Annotated Image
    for row in fd_ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        issues.append({
            'num': row[0],
            'checklist': '',
            'issue_id': row[1] or '',
            'screen': row[2] or '',
            'issue_type': row[3] or '',
            'wcag_nums': row[4] or '',
            'checkpoint_name': row[5] or '',
            'level': row[6] or 'A',
            'rationale': row[7] or '',
            'severity': row[8] or 'Moderate',
            'fix': '',
            'annotated_screenshot': row[9] or '',
        })
    wb.close()
    return screen_info, issues


def build_content_index(base_path):
    """Build content-based index from issue_index.json files for Format B matching.

    Also builds a screen-level fallback: if no issue-specific screenshot,
    use the general screen annotated screenshot for that checklist.
    """
    index = {}  # (issue_type, detail_prefix) -> annotated_screenshot_path
    screen_fallback = {}  # (screen_tag, checklist_folder) -> annotated_screenshot_path
    base = Path(base_path)
    if not base.exists():
        return index, screen_fallback

    for idx_file in sorted(base.glob("*/issues/issue_index.json")):
        checklist_folder = idx_file.parent.parent.name
        try:
            with open(idx_file) as f:
                data = json.load(f)
            for entry in data.get("issues", []):
                ann = entry.get("issue_annotated_screenshot", "")
                if not ann:
                    continue
                # Fix path: ann starts with "checklist_reports/XX_name/..."
                if ann.startswith("checklist_reports/"):
                    full_path = str(base / ann.replace("checklist_reports/", "", 1))
                else:
                    full_path = str(base / checklist_folder / ann)

                itype = entry.get("issue_type", "")
                detail = entry.get("detail", "")[:80]
                key = (itype, detail)
                if os.path.isfile(full_path):
                    index[key] = full_path

                # Build screen-level fallback
                screen_tag = entry.get("screen_tag", "")
                if screen_tag:
                    screen_ann = str(base / checklist_folder / "states" / "annotated" / f"state_{screen_tag}_annotated.png")
                    if os.path.isfile(screen_ann):
                        screen_fallback[(screen_tag, checklist_folder)] = screen_ann
        except Exception:
            continue

    # Also scan for screen-level annotated screenshots (without issue IDs)
    for ann_dir in base.glob("*/states/annotated"):
        checklist_folder = ann_dir.parent.parent.name
        for png in ann_dir.glob("*_annotated.png"):
            if "__ISSUE" not in png.name:
                # Extract screen_tag from filename like state_2026-03-27_174817_annotated.png
                tag = png.stem.replace("state_", "").replace("_annotated", "")
                screen_fallback[(tag, checklist_folder)] = str(png)

    print(f"  Content-based screenshot index: {len(index)} issue-level, {len(screen_fallback)} screen-level entries")
    return index, screen_fallback


def map_screenshots(issues, screenshot_index, content_index=None, screen_fallback=None):
    """Map each issue to its local screenshot path using filename, content, or screen fallback."""
    filename_hits = 0
    content_hits = 0
    fallback_hits = 0

    for issue in issues:
        ann = issue['annotated_screenshot']
        matched = False

        # Try filename match first
        if ann:
            lookup = ann if ann.endswith('.png') else ann + '_annotated.png'
            path = screenshot_index.get(lookup, '')
            if path:
                issue['local_screenshot'] = path
                matched = True
                filename_hits += 1

        # Fallback to content-based match (for Format B where filenames differ)
        if not matched and content_index:
            itype = issue.get('issue_type', '')
            detail = issue.get('rationale', '')[:80]
            key = (itype, detail)
            path = content_index.get(key, '')
            if path:
                issue['local_screenshot'] = path
                matched = True
                content_hits += 1

        # Screen-level fallback: use the annotated screenshot for the screen
        # Extract timestamp from the issue's annotated filename if available
        if not matched and screen_fallback and ann:
            # ann like "state_2026-03-23_010600__ISSUE-001_annotated.png"
            # extract timestamp between "state_" and "__ISSUE"
            import re
            m = re.search(r'state_(\d{4}-\d{2}-\d{2}_\d{6})', ann)
            if m:
                screen_tag = m.group(1)
                # Try any checklist folder with this screen timestamp
                for key, path in screen_fallback.items():
                    if key[0] == screen_tag:
                        issue['local_screenshot'] = path
                        matched = True
                        fallback_hits += 1
                        break

        if not matched:
            issue['local_screenshot'] = ''

    with_ss = sum(1 for i in issues if i['local_screenshot'])
    print(f"  Issues with screenshots: {with_ss}/{len(issues)} (filename:{filename_hits}, content:{content_hits}, screen-fallback:{fallback_hits})")


def list_existing_files(uploader):
    """List file names and IDs already in the OneDrive folder."""
    import requests
    try:
        if uploader._user_id == "me":
            base = "https://graph.microsoft.com/v1.0/me/drive"
        else:
            base = f"https://graph.microsoft.com/v1.0/users/{uploader._user_id}/drive"

        folder_path = uploader._folder.strip("/")
        url = f"{base}/root:/{folder_path}:/children?$top=1000&$select=name,id"

        existing = {}  # name -> item_id
        while url:
            resp = requests.get(url, headers=uploader._headers(), timeout=60)
            if resp.status_code == 404:
                break
            resp.raise_for_status()
            data = resp.json()
            for item in data.get("value", []):
                existing[item["name"]] = item["id"]
            url = data.get("@odata.nextLink")
            if len(existing) % 2000 == 0 and existing:
                print(f"    Indexed {len(existing)} existing files...")

        print(f"  Existing OneDrive files: {len(existing)}")
        return existing, base
    except Exception as e:
        print(f"  Warning: Could not list OneDrive files: {e}")
        return {}, ""


def batch_create_share_links(uploader, base_url, item_ids):
    """Create share links for multiple items using Graph batch API (20 per batch)."""
    import requests
    results = {}  # item_id -> share_url

    # Build user-scoped URL prefix for batch requests
    if uploader._user_id == "me":
        drive_prefix = "/me/drive"
    else:
        drive_prefix = f"/users/{uploader._user_id}/drive"

    total_batches = (len(item_ids) + 19) // 20
    batch_url = "https://graph.microsoft.com/v1.0/$batch"

    for batch_idx, batch_start in enumerate(range(0, len(item_ids), 20)):
        batch = item_ids[batch_start:batch_start + 20]
        requests_body = []
        for i, item_id in enumerate(batch):
            requests_body.append({
                "id": str(i),
                "method": "POST",
                "url": f"{drive_prefix}/items/{item_id}/createLink",
                "body": {"type": "view", "scope": "anonymous"},
                "headers": {"Content-Type": "application/json"},
            })

        for attempt in range(3):
            try:
                resp = requests.post(
                    batch_url,
                    headers={**uploader._headers(), "Content-Type": "application/json"},
                    json={"requests": requests_body},
                    timeout=60,
                )
                if resp.status_code == 429:
                    retry_after = int(resp.headers.get("Retry-After", 10))
                    print(f"    Rate limited, waiting {retry_after}s...")
                    time.sleep(retry_after)
                    continue
                resp.raise_for_status()
                data = resp.json()
                for response in data.get("responses", []):
                    idx = int(response["id"])
                    if response.get("status") in (200, 201):
                        link_url = response["body"]["link"]["webUrl"]
                        results[batch[idx]] = link_url
                break
            except Exception as e:
                if attempt == 2:
                    print(f"    Batch share link error: {e}")
                else:
                    time.sleep(2 ** attempt)

        if (batch_idx + 1) % 25 == 0:
            print(f"    Share links: {len(results)}/{len(item_ids)} ({batch_idx+1}/{total_batches} batches)")
        time.sleep(0.2)

    return results


def upload_screenshots(issues, uploader):
    """Upload all unique screenshots to OneDrive, skipping already-uploaded."""
    # Build a map of what's already on OneDrive
    existing, base_url = list_existing_files(uploader)

    upload_cache = {}
    unique_paths = sorted(set(
        i['local_screenshot'] for i in issues if i['local_screenshot']
    ))
    total = len(unique_paths)
    print(f"  Unique screenshots to upload: {total}")

    # Separate into skip vs upload
    to_skip = {}    # path -> item_id
    to_upload = []  # paths
    for path in unique_paths:
        fp = Path(path)
        filename = fp.name
        parts = fp.parts
        unique_prefix = ""
        for i, part in enumerate(parts):
            if part == "checklist_reports" and i + 1 < len(parts):
                unique_prefix = parts[i + 1] + "_"
                break
        upload_name = f"{unique_prefix}{filename}"

        if upload_name in existing:
            to_skip[path] = existing[upload_name]
        else:
            to_upload.append(path)

    print(f"  Already on OneDrive: {len(to_skip)}, Need upload: {len(to_upload)}")

    # Batch-create share links for existing files
    if to_skip:
        print(f"  Getting share links for {len(to_skip)} existing files (batch)...")
        item_ids = list(to_skip.values())
        path_list = list(to_skip.keys())
        share_links = batch_create_share_links(uploader, base_url, item_ids)
        for path, item_id in to_skip.items():
            upload_cache[path] = share_links.get(item_id, "")
        linked = sum(1 for v in share_links.values() if v)
        print(f"  Got {linked}/{len(to_skip)} share links")

    # Upload new files
    upload_count = 0
    upload_errors = 0
    for path in to_upload:
        for attempt in range(5):
            try:
                url = uploader.upload(path)
                upload_cache[path] = url
                upload_count += 1
                if upload_count % 25 == 0:
                    print(f"    Uploaded: {upload_count}/{len(to_upload)} ({100*(upload_count+len(to_skip))//total}% total)")
                time.sleep(0.3)
                break
            except Exception as e:
                err_str = str(e)
                if "429" in err_str or "Too Many Requests" in err_str:
                    wait = min(30, 5 * (attempt + 1))
                    print(f"    Rate limited, waiting {wait}s...")
                    time.sleep(wait)
                elif attempt == 4:
                    print(f"    FAILED: {os.path.basename(path)} - {e}")
                    upload_cache[path] = ""
                    upload_errors += 1
                else:
                    time.sleep(2 ** attempt)

    print(f"  Upload complete: {upload_count} new, {len(to_skip)} skipped, {upload_errors} failed")

    for issue in issues:
        path = issue.get('local_screenshot')
        issue['screenshot_url'] = upload_cache.get(path, '') if path else ''


def label_value(ws, row, label, value):
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = Font(name="Calibri", size=11)
    lc.border = _THIN
    vc = ws.cell(row=row, column=3, value=value)
    vc.font = Font(name="Calibri", size=11)
    vc.border = _THIN
    vc.alignment = Alignment(wrap_text=True)


def generate_report(app_config, screen_info, issues, out_file):
    """Generate the ICICI-format 3-sheet XLSX report."""
    app_name = app_config['name']
    app_version = app_config['version']

    wb = Workbook()

    # === COVER SHEET ===
    ws = wb.active
    ws.title = "Cover"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 25

    ws.merge_cells("A1:D3")
    c = ws.cell(row=1, column=1, value="Accessibility Audit Report")
    c.font = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
    c.fill = _HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(1, 4):
        ws.row_dimensions[r].height = 28
        for col in range(1, 5):
            ws.cell(row=r, column=col).fill = _HEADER_FILL

    ws.merge_cells("A5:D5")
    c5 = ws.cell(row=5, column=1, value="ASSET DETAILS")
    c5.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c5.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=5, column=col).fill = _SECTION_FILL

    label_value(ws, 6, "Asset / Product Name", app_name)
    label_value(ws, 7, "App Version", app_version)
    label_value(ws, 8, "Standard", "WCAG 2.1 A and AA")
    label_value(ws, 9, "Scan Mode", app_config.get('scan_mode', 'Android Mobile App'))
    label_value(ws, 10, "Screens Analyzed", str(len(screen_info)))

    ws.merge_cells("A12:D12")
    c12 = ws.cell(row=12, column=1, value="SUBMISSION DETAILS")
    c12.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c12.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=12, column=col).fill = _SECTION_FILL
    label_value(ws, 13, "Date of Submission", "14 Apr 2026")

    ws.merge_cells("A15:D15")
    c15 = ws.cell(row=15, column=1, value="ABOUT THIS AUDIT")
    c15.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c15.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=15, column=col).fill = _SECTION_FILL

    ws.merge_cells("A16:D16")
    c16 = ws.cell(row=16, column=1, value=(
        "This report presents findings from an accessibility audit conducted in accordance with the Web Content "
        "Accessibility Guidelines (WCAG) 2.1. Issues are categorised by severity (Blocker / High / Medium / Low) "
        "and mapped to the relevant WCAG success criteria. Conformance levels A and AA are referenced throughout."
    ))
    c16.font = Font(name="Calibri", size=10)
    c16.alignment = Alignment(wrap_text=True, vertical="top")
    c16.border = _THIN
    ws.row_dimensions[16].height = 72

    ws.merge_cells("A18:D18")
    c18 = ws.cell(row=18, column=1, value="AUDITOR DETAILS")
    c18.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c18.fill = _SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=18, column=col).fill = _SECTION_FILL
    label_value(ws, 19, "Auditor", "Automated WCAG Scanner")
    label_value(ws, 20, "Organisation", "OnFinance AI")

    ws.merge_cells("A22:D22")
    c22 = ws.cell(row=22, column=1, value="CONFIDENTIAL \u2014 For client use only")
    c22.font = Font(name="Calibri", size=11, italic=True)
    c22.alignment = Alignment(horizontal="center", vertical="center")
    c22.border = _THIN

    # === STATUS SHEET ===
    ws_s = wb.create_sheet("Status")
    for col_letter, w in zip("ABCDEFGH", [30, 45, 35, 18, 18, 18, 18, 18]):
        ws_s.column_dimensions[col_letter].width = w

    for col_idx, hdr in enumerate([
        "Screen Label", "Activity", "Notes", "Labels/Instructions",
        "Keyboard", "Focus Order", "Name, Role, Value", "Touch Targets",
    ], start=1):
        c = ws_s.cell(row=1, column=col_idx, value=hdr)
        c.fill = _SALMON_FILL
        c.font = _DARK_FONT
        c.alignment = _CENTER
        c.border = _THIN
    ws_s.freeze_panes = "A2"

    screen_checklists = {}
    screen_issue_types = {}
    for issue in issues:
        scr = issue['screen']
        screen_checklists.setdefault(scr, set()).add(issue['checklist'])
        screen_issue_types.setdefault(scr, set()).add(issue['issue_type'])

    for screen_label in sorted(screen_info.keys()):
        info = screen_info[screen_label]
        row_idx = ws_s.max_row + 1
        failed = screen_checklists.get(screen_label, set())

        for col_idx, val in enumerate([screen_label, info.get('activity', ''), ""], start=1):
            c = ws_s.cell(row=row_idx, column=col_idx, value=val)
            c.border = _THIN
            c.alignment = _WRAP_TOP

        for cl_name, col in [('Labels or Instructions', 4), ('Keyboard', 5),
                             ('Focus Order', 6), ('Name, Role, Value', 7)]:
            c = ws_s.cell(row=row_idx, column=col)
            if cl_name in failed:
                c.value = "Issues Found"
                c.fill = PatternFill("solid", fgColor="FFA500")
            else:
                c.value = "Done"
                c.fill = _GREEN_FILL
            c.border = _THIN
            c.alignment = _CENTER

        has_touch = 'small_touch_target' in screen_issue_types.get(screen_label, set())
        c8 = ws_s.cell(row=row_idx, column=8)
        c8.value = "Issues Found" if has_touch else "Done"
        c8.fill = PatternFill("solid", fgColor="FFA500") if has_touch else _GREEN_FILL
        c8.border = _THIN
        c8.alignment = _CENTER

    # === ISSUES SHEET ===
    ws_i = wb.create_sheet("Issues")
    col_widths = [30, 22, 18, 20, 12, 35, 45, 45, 20, 22, 45, 55, 14, 22, 12, 18]
    for i, w in enumerate(col_widths):
        ws_i.column_dimensions[get_column_letter(i + 1)].width = w

    headers_i = [
        "Issue Title", "Page Title", "Type", "AT/Browser", "Severity",
        "Action Performed", "Actual Result", "Expected Result",
        "Failed WCAG 2.1 checkpoint(s)", "Screencast",
        "Suggested Resolutions", "Suggested Code Fix",
        "Dev Status", "Dev Comments", "QA Status", "QA Comments",
    ]
    for col_idx, hdr in enumerate(headers_i, start=1):
        c = ws_i.cell(row=1, column=col_idx, value=hdr)
        c.fill = _SALMON_FILL
        c.font = _DARK_FONT
        c.alignment = _CENTER
        c.border = _THIN
    ws_i.freeze_panes = "A2"

    action_performed = (
        f"1) Launch the {app_name} App ({app_version}).\n"
        "2) Navigate to the target screen.\n"
        "3) Run automated WCAG 2.1 accessibility scan."
    )

    for idx, issue in enumerate(issues):
        it = ISSUE_TYPE_MAP.get(issue['issue_type'], 'Screen Reader')
        sev = SEVERITY_DISPLAY.get(issue['severity'], 'Medium')
        primary_wcag = issue['wcag_nums'].split(',')[0].strip() if issue['wcag_nums'] else ''

        expected = EXPECTED_RESULT.get(primary_wcag,
            f"Content must conform to WCAG 2.1 {primary_wcag} {issue['checkpoint_name']} (Level {issue['level']}).")

        what = WHAT_TO_FIX.get(primary_wcag, issue['fix'][:200] if issue['fix'] else '')
        how = HOW_TO_FIX.get(primary_wcag, '')
        suggested = (what + "\n\n" + "\n".join(how.split("\n")[:2])) if how else (issue['fix'] or what)

        code_fix = CODE_FIX.get(primary_wcag, '')
        title_summary = issue['rationale'][:80] if issue['rationale'] else issue['issue_type']
        issue_title = f"{it}: {issue['screen']}: {title_summary}"
        wcag_str = f"{issue['wcag_nums']} {issue['checkpoint_name']} (Level {issue['level']})"

        row_idx = ws_i.max_row + 1
        ws_i.row_dimensions[row_idx].height = 80

        row_values = [
            issue_title, issue['screen'], it,
            "Android (Automated WCAG Scanner)", sev, action_performed,
            issue['rationale'], expected, wcag_str,
            "", suggested, code_fix, "", "", "", "",
        ]
        for col_idx, val in enumerate(row_values, start=1):
            c = ws_i.cell(row=row_idx, column=col_idx, value=val)
            c.border = _THIN
            c.alignment = _WRAP_TOP

        sev_style = _SEV_FILLS.get(sev)
        if sev_style:
            ws_i.cell(row=row_idx, column=5).fill = sev_style[0]
            ws_i.cell(row=row_idx, column=5).font = sev_style[1]

        sc_cell = ws_i.cell(row=row_idx, column=10)
        if issue.get('screenshot_url'):
            sc_cell.value = "View Screenshot"
            sc_cell.hyperlink = issue['screenshot_url']
            sc_cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            sc_cell.alignment = Alignment(horizontal="center", vertical="center")
        elif issue.get('annotated_screenshot'):
            sc_cell.value = issue['annotated_screenshot']
            sc_cell.alignment = _WRAP_TOP

        if code_fix:
            code_cell = ws_i.cell(row=row_idx, column=12)
            code_cell.font = _CODE_FIX_FONT
            code_cell.fill = _CODE_FIX_FILL
            code_cell.alignment = Alignment(wrap_text=True, vertical="top")

    print(f"  Issues written: {len(issues)}")
    wb.save(out_file)
    print(f"  Report saved: {out_file} ({wb.sheetnames})")


def process_app(app_config, uploader):
    """Full pipeline for one app: read → map screenshots → upload → generate."""
    name = app_config['name']
    print(f"\n{'='*60}")
    print(f"Processing: {name}")
    print(f"{'='*60}")

    # 1. Read source
    if app_config['format'] == 'A':
        screen_info, issues = read_source_format_a(app_config['src_file'])
    else:
        screen_info, issues = read_source_format_b(app_config['src_file'])
    print(f"  Screens: {len(screen_info)}, Issues: {len(issues)}")

    # 2. Build screenshot index and map
    ss_index = build_screenshot_index(app_config['base'])
    content_idx = None
    screen_fb = None
    if app_config['format'] == 'B':
        content_idx, screen_fb = build_content_index(app_config['base'])
    map_screenshots(issues, ss_index, content_idx, screen_fb)

    # 3. Upload to OneDrive
    print("  Uploading screenshots to OneDrive...")
    upload_screenshots(issues, uploader)

    # 4. Generate report
    print("  Generating ICICI-format report...")
    generate_report(app_config, screen_info, issues, app_config['out_file'])
    print(f"  Done: {name}")


def main():
    sys.stdout.reconfigure(line_buffering=True)

    from src.accessibility_scanner.image_uploader import OneDriveUploader
    uploader = OneDriveUploader()

    for app in APPS:
        process_app(app, uploader)

    print(f"\n{'='*60}")
    print("All reports complete!")
    for app in APPS:
        print(f"  {app['name']}: {app['out_file']}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
